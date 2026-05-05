from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from connections import connect_to_epm_environment
from generate_dimension_reports import (
    build_report_rows,
    discover_dimensions_with_links,
    export_data,
    export_data_slice,
    fetch_dimension_alias_map,
    fetch_dimension_members,
    fetch_l0_descendants_under_parent,
    get_last_dataexport_failures,
    write_report,
)


BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "reports" / "EPM_Actual_vs_Forecast_Insight_Output.xlsx"
DEFAULT_OUTPUT = BASE_DIR / "reports" / "EPM_Actual_vs_Forecast_Intersection_Output_UI.xlsx"
YEAR_OPTIONS = ["FY23", "FY24", "FY25", "FY26", "FY27", "FY28", "FY29"]
PERIOD_OPTIONS = [f"FSC{i}" for i in range(1, 13)]
SCENARIO_OPTIONS = ["Actual", "Forecast"]
ACTUAL_VERSION_OPTIONS = ["Sample_Actual", "Final", "Working", "Base"]
FORECAST_VERSION_OPTIONS = ["Sample_Plan", "Working", "Final", "Base", "Sample_Actual"]
ENTITY_OPTIONS = ["LE_ANZ", "LE_USA", "LE_UK", "LE_IND"]
POV_LOCATION_OPTIONS = ["Total_Location_Input", "Total Location"]
POV_UOM_OPTIONS = ["No UOM"]
POV_LOB_OPTIONS = ["No LOB"]
POV_ENERGY_TYPE_OPTIONS = ["Pollutant 1", "All Pollutants"]


def _resolve_output_path(raw_path: str) -> Path:
    s = str(raw_path or "").strip()
    if not s:
        p = DEFAULT_OUTPUT
    else:
        p = Path(s)
        if not p.is_absolute():
            p = BASE_DIR / p
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def _validate_template_path(template_path: Path) -> None:
    if not template_path.exists():
        raise FileNotFoundError(f"Template file not found: {template_path}")


def _split_csv(s: str) -> list[str]:
    return [x.strip() for x in s.split(",") if x.strip()]


def _parse_year_int(v: str) -> int | None:
    s = str(v or "").strip()
    if not s:
        return None
    digits = "".join(ch for ch in s if ch.isdigit())
    if not digits:
        return None
    try:
        if len(digits) == 2:
            return 2000 + int(digits)
        if len(digits) >= 4:
            return int(digits[-4:])
        return int(digits)
    except Exception:
        return None


def _to_fy_label(year_int: int, sample: str | None = None) -> str:
    s = str(sample or "").strip()
    prefix = "FY" if s.upper().startswith("FY") else ""
    yy = year_int % 100
    return f"{prefix}{yy:02d}" if prefix else str(year_int)


def _expand_years_with_two_priors(years: list[str]) -> list[str]:
    out: list[str] = list(years)
    seen = {_parse_year_int(y): y for y in years if _parse_year_int(y) is not None}
    for y in years:
        yi = _parse_year_int(y)
        if yi is None:
            continue
        for prior in (yi - 1, yi - 2):
            if prior in seen:
                continue
            label = _to_fy_label(prior, y)
            out.append(label)
            seen[prior] = label
    return out


def _pick_field_name(items: list[dict[str, Any]], candidates: list[str]) -> str:
    if not items:
        return ""
    lower_map: dict[str, str] = {}
    for row in items:
        if not isinstance(row, dict):
            continue
        for k in row.keys():
            kl = str(k).lower()
            if kl not in lower_map:
                lower_map[kl] = str(k)
    for c in candidates:
        if c.lower() in lower_map:
            return lower_map[c.lower()]
    return ""


@st.cache_data(show_spinner=False, ttl=1800)
def _load_entity_members_and_aliases() -> tuple[list[str], dict[str, str]]:
    try:
        session, cfg = connect_to_epm_environment()
        cfg["plantype"] = cfg.get("plantype") or cfg["application"]
        dimensions, dim_links = discover_dimensions_with_links(session, cfg)
        dim_lookup = {d.lower(): d for d in dimensions}
        entity_dim = dim_lookup.get("entity")
        if not entity_dim:
            return ENTITY_OPTIONS, {}
        members = fetch_dimension_members(session, cfg, entity_dim, dim_links.get(entity_dim, ""))
        aliases = fetch_dimension_alias_map(session, cfg, entity_dim, dim_links.get(entity_dim, ""))
        return (members or ENTITY_OPTIONS), aliases
    except Exception:
        return ENTITY_OPTIONS, {}


def _read_data_sheet_as_df(path: Path) -> pd.DataFrame:
    try:
        wb = load_workbook(path, data_only=True)
        if "Data" not in wb.sheetnames:
            return pd.DataFrame()
        ws = wb["Data"]
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        return pd.DataFrame()
    if not rows:
        return pd.DataFrame()
    headers = list(rows[0])
    data = [r for r in rows[1:] if any(v is not None and str(v).strip() != "" for v in r)]
    return pd.DataFrame(data, columns=headers)


def _read_insights_sheet_as_df(path: Path) -> pd.DataFrame:
    try:
        wb = load_workbook(path, data_only=True)
        if "Insights" not in wb.sheetnames:
            return pd.DataFrame()
        ws = wb["Insights"]
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        return pd.DataFrame()
    if not rows:
        return pd.DataFrame()

    header_idx = None
    headers: list[Any] = []
    for i, r in enumerate(rows):
        vals = [str(v).strip() if v is not None else "" for v in r]
        if "Entity" in vals and "Account" in vals:
            header_idx = i
            headers = list(r)
            break

    if header_idx is None:
        return pd.DataFrame()

    data_rows = []
    for r in rows[header_idx + 1 :]:
        if any(v is not None and str(v).strip() != "" for v in r):
            data_rows.append(r)

    if not data_rows:
        return pd.DataFrame(columns=headers)

    canonical = [
        "Entity",
        "Account",
        "Actual",
        "Forecast",
        "Prior",
        "Prior to Prior",
        "Prior 1",
        "Prior 2",
        "Variance",
        "Variance(%)",
        "Var %",
        "Trend",
        "Pattern",
        "Severity",
        "Narrative Message",
        "Recommended Action",
    ]

    first_index: dict[str, int] = {}
    for i, h in enumerate(headers):
        key = str(h).strip()
        if key and key not in first_index:
            first_index[key] = i

    selected_cols = [c for c in canonical if c in first_index]
    if not selected_cols:
        selected_cols = list(first_index.keys())

    selected_idx = [first_index[c] for c in selected_cols]
    records = []
    for r in data_rows:
        row_vals = list(r)
        rec = []
        for idx in selected_idx:
            rec.append(row_vals[idx] if idx < len(row_vals) else None)
        records.append(rec)

    return pd.DataFrame(records, columns=selected_cols)


def _to_numeric_series(df: pd.DataFrame, col: str) -> pd.Series:
    if col not in df.columns:
        return pd.Series([0.0] * len(df), index=df.index, dtype="float64")
    return pd.to_numeric(df[col], errors="coerce").fillna(0.0)


def _pick_existing_column(df: pd.DataFrame, candidates: list[str]) -> str:
    for c in candidates:
        if c in df.columns:
            return c
    return ""


def _friendly_member_name(v: Any) -> str:
    s = str(v or "").strip()
    if "/" in s:
        parts = [p for p in s.split("/") if p]
        if parts:
            return parts[-1]
    return s


def _format_display_df(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "Entity" in out.columns:
        out = out.drop(columns=["Entity"])
    for name_col in ("Entity", "Account"):
        if name_col in out.columns:
            out[name_col] = out[name_col].map(_friendly_member_name)

    currency_cols = ["Actual", "Forecast", "Prior", "Prior to Prior", "Prior 1", "Prior 2", "Variance"]
    for c in currency_cols:
        if c in out.columns:
            out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0.0).map(lambda v: f"{v:,.2f}$")

    pct_col = _pick_existing_column(out, ["Variance(%)", "Var %"])
    if pct_col:
        out[pct_col] = pd.to_numeric(out[pct_col], errors="coerce").fillna(0.0).map(lambda v: f"{v:.2f}%")
    return out


def _build_summary_from_df(df: pd.DataFrame) -> dict[str, Any]:
    actual = _to_numeric_series(df, "Actual")
    forecast = _to_numeric_series(df, "Forecast")
    variance = _to_numeric_series(df, "Variance")

    total_actual = float(actual.sum())
    total_forecast = float(forecast.sum())
    total_variance = float(variance.sum())
    variance_pct = (total_variance / total_forecast) if total_forecast else 0.0

    pattern_series = df["Pattern"].astype(str).str.strip() if "Pattern" in df.columns else pd.Series([], dtype="object")
    severity_series = df["Severity"].astype(str).str.strip() if "Severity" in df.columns else pd.Series([], dtype="object")

    return {
        "total_actual": total_actual,
        "total_forecast": total_forecast,
        "total_variance": total_variance,
        "variance_pct": variance_pct,
        "pattern_counts": {
            "Sustained upside": int((pattern_series == "Sustained upside").sum()),
            "Sustained downside": int((pattern_series == "Sustained downside").sum()),
            "Forecast bias": int((pattern_series == "Forecast bias").sum()),
            "Outlier": int((pattern_series == "Outlier").sum()),
        },
        "severity_counts": {
            "High": int((severity_series == "High").sum()),
            "Medium": int((severity_series == "Medium").sum()),
            "Low": int((severity_series == "Low").sum()),
        },
    }


def run_intersection_report(
    entity_parent: str,
    measure_parent: str,
    period_parent: str,
    entity_members_selected: list[str],
    account_members_selected: list[str],
    periods_selected: list[str],
    years: list[str],
    scenarios: list[str],
    versions: list[str],
    forecast_versions: list[str],
    pov_overrides: dict[str, str | list[str]] | None,
    output_path: Path,
) -> tuple[Path, int, list[str]]:
    _validate_template_path(TEMPLATE_PATH)
    session, cfg = connect_to_epm_environment()
    cfg["plantype"] = cfg.get("plantype") or cfg["application"]

    dimensions, dim_links = discover_dimensions_with_links(session, cfg)
    dim_lookup = {d.lower(): d for d in dimensions}

    entity_dim = dim_lookup.get("entity")
    account_dim = dim_lookup.get("account") or dim_lookup.get("accounts") or dim_lookup.get("measure") or dim_lookup.get("measures")
    period_dim = dim_lookup.get("period")

    if not entity_dim or not account_dim:
        raise RuntimeError("Entity or Account/Measures dimension not found.")

    entities = fetch_l0_descendants_under_parent(session, cfg, entity_dim, dim_links.get(entity_dim, ""), [entity_parent])
    measures = fetch_l0_descendants_under_parent(session, cfg, account_dim, dim_links.get(account_dim, ""), [measure_parent])
    if entity_members_selected:
        entities = entity_members_selected
    if account_members_selected:
        measures = account_members_selected

    periods: list[str] = []
    if periods_selected:
        periods = periods_selected
    elif period_dim:
        periods = fetch_l0_descendants_under_parent(session, cfg, period_dim, dim_links.get(period_dim, ""), [period_parent])
        if not periods and period_parent:
            all_periods = fetch_dimension_members(session, cfg, period_dim, dim_links.get(period_dim, ""))
            periods = [period_parent] if any(p.lower() == period_parent.lower() for p in all_periods) else []

    if not entities:
        entities = fetch_dimension_members(session, cfg, entity_dim, dim_links.get(entity_dim, ""))
    if not measures:
        measures = fetch_dimension_members(session, cfg, account_dim, dim_links.get(account_dim, ""))
    entity_alias_map = fetch_dimension_alias_map(session, cfg, entity_dim, dim_links.get(entity_dim, ""))
    account_alias_map = fetch_dimension_alias_map(session, cfg, account_dim, dim_links.get(account_dim, ""))

    actual_versions = versions if versions else ["Sample_Actual"]
    forecast_version_candidates = forecast_versions if forecast_versions else ["Sample_Plan"]
    forecast_years = _expand_years_with_two_priors(years if years else ["FY26"])

    # Fetch Actual and Forecast through the same slice API path for consistent scenario retrieval.
    actual_items = export_data_slice(
        session,
        cfg,
        dimensions,
        dim_links,
        entity_dim,
        account_dim,
        entities,
        measures,
        ["Actual"],
        years if years else ["FY26"],
        periods if periods else ["FSC10"],
        actual_versions,
        pov_overrides=pov_overrides,
    )
    forecast_items = export_data_slice(
        session,
        cfg,
        dimensions,
        dim_links,
        entity_dim,
        account_dim,
        entities,
        measures,
        ["Forecast"],
        forecast_years,
        periods if periods else ["FSC10"],
        forecast_version_candidates,
        pov_overrides=pov_overrides,
    )
    items = actual_items + forecast_items

    # Fall back to dataexport only when slice APIs return no rows.
    if not items:
        items = export_data(session, cfg, dimensions)
    entity_field = _pick_field_name(items, ["Entity"]) or "Entity"
    account_field = _pick_field_name(items, ["Account", "Measures", "Measure"]) or "Account"
    scenario_field = _pick_field_name(items, ["Scenario"]) or "Scenario"
    data_field = _pick_field_name(items, ["Data", "Value", "Amount"]) or "Data"

    version_field = _pick_field_name(items, ["Version"])
    has_version_field = bool(version_field)

    rows = build_report_rows(
        items,
        entities,
        measures,
        entity_field,
        account_field,
        scenario_field,
        data_field,
        include_all_intersections=not bool(items),
        allowed_scenarios={s.lower() for s in scenarios} if scenarios else None,
        allowed_periods={p.lower() for p in periods} if periods else None,
        allowed_years={y.lower() for y in forecast_years} if forecast_years else None,
        allowed_versions=(
            {v.lower() for v in (actual_versions + forecast_version_candidates)}
            if has_version_field
            else None
        ),
        current_year=years[0] if years else "FY26",
        entity_alias_map=entity_alias_map,
        account_alias_map=account_alias_map,
    )

    notes = [
        f"Entity L0 descendants under parent '{entity_parent}': {len(entities)}",
        f"{account_dim} L0 descendants under parent '{measure_parent}': {len(measures)}",
        f"Periods used: {periods}",
        f"Scenarios filter: {scenarios}",
        f"Years filter: {years}",
        f"Forecast years fetched (incl priors): {forecast_years}",
        f"Actual version filter: {versions}",
        f"Forecast version filter: {forecast_version_candidates}",
        f"POV overrides: {pov_overrides or {}}",
    ]

    output_path = _resolve_output_path(str(output_path))
    final_output = write_report(TEMPLATE_PATH, output_path, rows, dimensions, notes)
    return final_output, len(rows), dimensions


st.set_page_config(page_title="EPM Data Analysis", layout="wide")
logo_col, title_col = st.columns([1, 6])
with logo_col:
    st.image("https://upload.wikimedia.org/wikipedia/commons/5/50/Oracle_logo.svg", width=120)
with title_col:
    title_placeholder = st.empty()
    title_placeholder.title("EPM Data Analysis")
st.markdown(
    """
    <style>
    :root {
        --rw-bg-1: #f7f8fa;
        --rw-bg-2: #eef1f5;
        --rw-surface: #ffffff;
        --rw-text: #1f2937;
        --rw-muted: #5b6472;
        --rw-border: #d6dbe4;
        --rw-accent: #c74634;
        --rw-accent-dark: #a33829;
        --rw-accent-soft: #fde9e5;
        --rw-focus: #86a8ff;
    }
    .stApp {
        background: linear-gradient(180deg, var(--rw-bg-1) 0%, var(--rw-bg-2) 100%);
        color: var(--rw-text);
        font-family: "Oracle Sans", "Segoe UI", "Helvetica Neue", Arial, sans-serif;
    }
    h1, h2, h3 {
        color: var(--rw-text);
        letter-spacing: 0.2px;
    }
    [data-testid="stTabs"] {
        background: var(--rw-surface);
        border: 1px solid var(--rw-border);
        border-radius: 10px;
        padding: 0.25rem 0.35rem;
        box-shadow: 0 1px 2px rgba(31, 41, 55, 0.05);
    }
    [data-testid="stTabs"] button {
        border-radius: 8px;
        color: var(--rw-muted);
    }
    [data-testid="stTabs"] button[aria-selected="true"] {
        background: var(--rw-accent-soft) !important;
        color: var(--rw-accent) !important;
        font-weight: 600;
    }
    div[data-testid="stMetricValue"] {
        color: var(--rw-text);
    }
    div[data-testid="stMetric"] {
        background: var(--rw-surface);
        border: 1px solid var(--rw-border);
        border-radius: 10px;
        padding: 0.6rem 0.8rem;
        box-shadow: 0 1px 2px rgba(31, 41, 55, 0.04);
    }
    div.stButton > button {
        background: var(--rw-accent);
        color: white;
        border: 1px solid var(--rw-accent);
        border-radius: 8px;
        padding: 0.5rem 1rem;
        font-weight: 600;
        box-shadow: 0 2px 5px rgba(199, 70, 52, 0.25);
    }
    div.stButton > button:hover {
        background: var(--rw-accent-dark);
        border-color: var(--rw-accent-dark);
        color: white;
    }
    div.stButton > button:focus {
        outline: 2px solid var(--rw-focus);
        outline-offset: 1px;
    }
    [data-baseweb="select"] > div,
    [data-testid="stTextInput"] input {
        border-radius: 8px !important;
        border: 1px solid var(--rw-border) !important;
        background: var(--rw-surface) !important;
        color: var(--rw-text) !important;
    }
    [data-baseweb="select"] > div:focus-within,
    [data-testid="stTextInput"] input:focus {
        border-color: var(--rw-focus) !important;
        box-shadow: 0 0 0 1px var(--rw-focus) !important;
    }
    [data-testid="stDataFrame"] {
        border-radius: 10px;
        overflow: hidden;
        border: 1px solid var(--rw-border);
        background: var(--rw-surface);
    }
    </style>
    """,
    unsafe_allow_html=True,
)

tab1, tab2 = st.tabs(["EPM Data Analysis", "View dimension members"])

with tab2:
    st.subheader("View dimension members")
    if st.button("Load dimensions and members"):
        try:
            with st.spinner("Connecting to EPM and fetching metadata..."):
                session, cfg = connect_to_epm_environment()
                cfg["plantype"] = cfg.get("plantype") or cfg["application"]
                dimensions, dim_links = discover_dimensions_with_links(session, cfg)
                st.success(f"Loaded {len(dimensions)} dimensions")
                for d in dimensions:
                    with st.expander(d, expanded=False):
                        members = fetch_dimension_members(session, cfg, d, dim_links.get(d, ""))
                        st.write(f"Member count: {len(members)}")
                        st.dataframe(pd.DataFrame({"member": members[:500]}), width="stretch")
        except Exception as e:
            st.error(f"Failed to load dimensions: {e}")

with tab1:
    st.subheader("Variance Insight Workspace")
    entity_option_values, entity_alias_map_ui = _load_entity_members_and_aliases()
    # Defensive normalization for Streamlit multiselect to avoid UI/runtime issues
    # when metadata APIs return null/duplicate/non-string member values.
    entity_option_values = [str(v).strip() for v in entity_option_values if str(v).strip()]
    entity_option_values = list(dict.fromkeys(entity_option_values))
    entity_alias_map_ui = {
        str(k).strip(): str(v).strip()
        for k, v in (entity_alias_map_ui or {}).items()
        if str(k).strip()
    }
    default_entity_selection = ["LE_ANZ"] if "LE_ANZ" in entity_option_values else (entity_option_values[:1] if entity_option_values else [])
    col1, col2 = st.columns(2)
    with col1:
        entity_members_selected = st.multiselect(
            "Entity (Default Alias)",
            options=entity_option_values,
            default=default_entity_selection,
            format_func=lambda m: entity_alias_map_ui.get(str(m), str(m)),
            key="entity_members_selected",
        )

    selected_entity_labels = [entity_alias_map_ui.get(m, m) for m in entity_members_selected if str(m).strip()]
    if selected_entity_labels:
        title_placeholder.title(f"EPM Data Analysis - {', '.join(selected_entity_labels)}")
    else:
        title_placeholder.title("EPM Data Analysis")
    with col2:
        # Period selection hidden from UI as requested; use all periods by default.
        periods_selected = PERIOD_OPTIONS
        years_selected = st.multiselect("Years", options=YEAR_OPTIONS, default=["FY26"])
        # Hidden defaults (removed from UI by request)
        scenarios_selected = ["Actual", "Forecast"]
        versions_selected = ["Sample_Actual"]
        forecast_versions_selected = ["Sample_Plan"]
        location_pov = "Total_Location_Input"
        uom_pov = "No UOM"
        lob_pov = "No LOB"
        energy_type_pov = "Pollutant 1"

    with st.expander("Advanced Filters", expanded=False):
        a1, a2 = st.columns(2)
        with a1:
            entity_parent = st.text_input("Entity parent", value="DO NOT USE")
            measure_parent = st.text_input("Measure parent", value="Profit and Loss Statement")
            period_parent = "YearTotal"
        with a2:
            entity_members_custom = st.text_input("Extra entity members (comma-separated, optional)", value="")
            account_members_csv = st.text_input("Account members (comma-separated, optional)", value="")
            output_file = st.text_input("Output file", value=str(DEFAULT_OUTPUT))

    if st.button("Generate intersection report"):
        out: Path | None = None
        try:
            with st.spinner("Generating workbook..."):
                entity_members_final = list(dict.fromkeys(entity_members_selected + _split_csv(entity_members_custom)))
                pov_overrides: dict[str, str | list[str]] = {}
                if location_pov.strip():
                    pov_overrides["Location"] = location_pov.strip()
                if uom_pov.strip():
                    pov_overrides["UoM"] = uom_pov.strip()
                if lob_pov.strip():
                    pov_overrides["LOB"] = lob_pov.strip()
                if energy_type_pov.strip():
                    pov_overrides["Energy Type"] = energy_type_pov.strip()

                out, count, dims = run_intersection_report(
                    entity_parent=entity_parent,
                    measure_parent=measure_parent,
                    period_parent=period_parent,
                    entity_members_selected=entity_members_final,
                    account_members_selected=_split_csv(account_members_csv),
                    periods_selected=periods_selected,
                    years=years_selected,
                    scenarios=scenarios_selected,
                    versions=versions_selected,
                    forecast_versions=forecast_versions_selected,
                    pov_overrides=pov_overrides,
                    output_path=_resolve_output_path(output_file),
                )
        except Exception as e:
            st.error(f"Report generation failed: {e}")
            st.info("Please verify EPM REST endpoint/app credentials and try again.")
            st.stop()

        if out is None or not Path(out).exists():
            st.error("Report generation did not produce an output file.")
            st.info("No file available to download. Please review connection/filters and retry.")
            st.stop()

        st.success("Report generated successfully.")
        with open(out, "rb") as f:
            st.download_button(
                label="Download Excel Report",
                data=f.read(),
                file_name=Path(out).name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        df = _read_data_sheet_as_df(out)
        if df.empty:
            failures = get_last_dataexport_failures()
            if failures:
                st.warning(
                    "No application rows were returned from the configured EPM data endpoints. "
                    "Endpoint responses indicate access/endpoint restrictions."
                )
                st.caption("Last endpoint attempts (status -> URL):")
                diag = pd.DataFrame(
                    [{"status": str(status), "url": url} for url, status in failures],
                    columns=["status", "url"],
                )
                st.dataframe(diag, width="stretch")

        if not df.empty:
            summary = _build_summary_from_df(df)
            st.markdown("### Summary")
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Total Actual", f"{summary['total_actual']:,.2f}$")
            m2.metric("Total Forecast", f"{summary['total_forecast']:,.2f}$")
            m3.metric("Total Variance", f"{summary['total_variance']:,.2f}$")
            m4.metric("Variance %", f"{summary['variance_pct']:.2%}")

            st.markdown("#### Top Patterns")
            pattern_df = pd.DataFrame(
                {
                    "Pattern": list(summary["pattern_counts"].keys()),
                    "Count": list(summary["pattern_counts"].values()),
                }
            )
            st.bar_chart(pattern_df.set_index("Pattern")["Count"])

        st.markdown("### Data preview")
        st.dataframe(_format_display_df(df).head(1000), width="stretch")

        insights_df = _read_insights_sheet_as_df(out)
        st.markdown("### Insights")
        if insights_df.empty:
            st.info("No rows found in Insights sheet.")
            st.caption("Trend line graph appears only when Insights rows with a Trend column are available.")
        else:
            st.caption("Business-friendly narrative view for stakeholders.")

            st.markdown("#### Trend Analysis (Trend column)")
            trend_col = _pick_existing_column(insights_df, ["Trend"])
            if trend_col:
                trend_series = insights_df[trend_col].astype(str).str.strip()
                trend_labels = ["Uptrend", "Downtrend", "Flat/volatile"]
                trend_line_df = pd.DataFrame({"Point": range(1, len(trend_series) + 1)})
                for label in trend_labels:
                    trend_line_df[label] = (trend_series == label).astype(int).cumsum()
                st.line_chart(trend_line_df.set_index("Point"), width="stretch")
            else:
                st.info("Trend column not found in Insights data, so line graph cannot be displayed.")

            st.markdown("#### Actual vs Forecast vs Prior 1 vs Prior 2")
            actual_col = _pick_existing_column(insights_df, ["Actual"])
            forecast_col = _pick_existing_column(insights_df, ["Forecast"])
            prior1_col = _pick_existing_column(insights_df, ["Prior 1", "Prior"])
            prior2_col = _pick_existing_column(insights_df, ["Prior 2", "Prior to Prior"])

            if actual_col and forecast_col and prior1_col and prior2_col:
                line_view = insights_df.copy()
                x_label_col = _pick_existing_column(line_view, ["Account"])
                if x_label_col:
                    line_view["Point"] = line_view[x_label_col].astype(str).str.strip().replace("", pd.NA).fillna(
                        pd.Series(range(1, len(line_view) + 1), index=line_view.index).map(lambda n: f"Point {n}")
                    )
                else:
                    line_view["Point"] = pd.Series(range(1, len(line_view) + 1), index=line_view.index).map(lambda n: f"Point {n}")

                metrics_line_df = pd.DataFrame(
                    {
                        "Point": line_view["Point"],
                        "Actual": pd.to_numeric(line_view[actual_col], errors="coerce").fillna(0.0),
                        "Forecast": pd.to_numeric(line_view[forecast_col], errors="coerce").fillna(0.0),
                        "Prior 1": pd.to_numeric(line_view[prior1_col], errors="coerce").fillna(0.0),
                        "Prior 2": pd.to_numeric(line_view[prior2_col], errors="coerce").fillna(0.0),
                    }
                )
                st.line_chart(metrics_line_df.set_index("Point"), width="stretch")
            else:
                st.info("Required columns for the 4-line chart were not found (Actual, Forecast, Prior 1, Prior 2).")

            insights_display_df = insights_df.drop(columns=["Entity"], errors="ignore")
            st.dataframe(insights_display_df.head(1000), width="stretch")

            if "Narrative Message" in insights_df.columns:
                st.markdown("#### Narrative Highlights")
                nar_cols = [c for c in ["Account", "Severity", "Narrative Message", "Recommended Action"] if c in insights_df.columns]
                st.dataframe(insights_df[nar_cols].head(20), width="stretch")

        var_pct_col = _pick_existing_column(df, ["Variance(%)", "Var %"])
        if not df.empty and var_pct_col:
            top = df.sort_values(by=var_pct_col, key=lambda s: s.abs(), ascending=False).head(5)
            st.markdown("### Top 5 by |Variance(%)|")
            st.bar_chart(top.set_index("Account")[var_pct_col])
