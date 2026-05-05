"""Validation script that cross-checks Actuals measures and emails results."""

from __future__ import annotations

import argparse
import base64
import csv
import json
import os
import re
import shutil
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Dict, Tuple
from urllib.parse import quote

from docx import Document
from docx.shared import Pt, RGBColor
import requests
from dotenv import load_dotenv, dotenv_values


SCRIPT_DIR = Path(__file__).resolve().parent


def _normalize_base_url(base_url: str) -> str:
    """Normalize EPM base URL to host root for Interop/Planning REST APIs."""
    base = (base_url or "").strip().rstrip("/")
    if not base:
        return base

    for suffix in ("/epmcloud/HyperionPlanning", "/HyperionPlanning", "/epmcloud"):
        if base.endswith(suffix):
            base = base[: -len(suffix)]
    return base.rstrip("/")


def _load_config() -> Dict[str, str]:
    primary_env = SCRIPT_DIR / ".env"
    fallback_env = SCRIPT_DIR / "env (3)"

    load_dotenv(primary_env, override=True)
    fallback_values = dotenv_values(fallback_env) if fallback_env.exists() else {}

    def _env(key: str, default: str = "") -> str:
        # Prefer active process/.env value; fall back to "env (3)" when missing.
        value = (os.getenv(key, "") or "").strip()
        if value:
            return value
        return (str(fallback_values.get(key, default)) or "").strip()

    cfg = {
        "base_url": _normalize_base_url(_env("EPM_BASE_URL", "")),
        "username": _env("EPM_USERNAME", ""),
        "password": _env("EPM_PASSWORD", ""),
        "api_version": _env("EPM_API_VERSION", "v3"),
        "application": _env("EPM_APPLICATION", ""),
        "admin_email": _env("ADMIN_EMAIL_ID", ""),
        "admin_name": _env("ADMIN_NAME", ""),
        "local_folder": _env("LOCAL_FOLDER", "") or os.getcwd(),
        "upload_version": _env("EPM_UPLOAD_VERSION", "") or _env("EPM_INTEROP_VERSION", "11.1.2.3.600"),
    }

    # Only core runtime fields are mandatory for validation.
    # Email/upload fields are optional unless notification is requested.
    required = ["base_url", "username", "password", "api_version", "application", "local_folder"]
    missing = [k for k in required if not cfg.get(k)]
    if missing:
        raise ValueError(f"Missing required environment variables: {', '.join(missing)}")

    return cfg


def _find_latest_export_zip(local_folder: Path) -> Path:
    candidates = sorted(
        local_folder.glob("*Export*Account*.zip"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError("No 'Export Account' zip found in the folder")

    for candidate in candidates:
        if zipfile.is_zipfile(candidate):
            return candidate

    raise zipfile.BadZipFile("No valid ZIP file found among '*Export*Account*.zip' candidates")


def _extract_measures_csv(zip_path: Path) -> Path:
    temp_dir = Path(tempfile.mkdtemp(prefix="validation_zip_"))
    with zipfile.ZipFile(zip_path, "r") as zf:
        target_member = None
        for member in zf.namelist():
            if re.search(r"Measures\.csv$", member, re.IGNORECASE):
                target_member = member
                break
        if not target_member:
            raise FileNotFoundError("No Measures.csv file found inside the zip")
        zf.extract(target_member, temp_dir)
    return temp_dir / target_member


def _find_local_measures_csv(local_folder: Path) -> Path:
    candidates = sorted(
        local_folder.glob("*ExportedMetadata*Measures*.csv"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError("No local exported Measures CSV found")
    return candidates[0]


def _load_measures(measures_csv: Path) -> set[str]:
    measures = set()
    with open(measures_csv, newline="", encoding="utf-8-sig") as csv_file:
        reader = csv.DictReader(csv_file)
        fieldnames = reader.fieldnames or []
        if "Measures" not in fieldnames:
            raise ValueError("Measures column not found in metadata CSV")
        for row in reader:
            value = (row.get("Measures") or "").strip()
            if value:
                measures.add(value)
    return measures


def _load_actual_measures(actuals_csv: Path) -> Tuple[int, list[str]]:
    with open(actuals_csv, newline="", encoding="utf-8-sig") as csv_file:
        reader = csv.reader(csv_file)
        rows = list(reader)
        if not rows:
            return 0, []
        header = rows[0]
        if len(header) < 9:
            raise ValueError("Actuals file does not have at least 9 columns")
        values = []
        for row in rows[1:]:
            if len(row) >= 9:
                values.append(row[8].strip())
        return len(rows) - 1, values


def _build_validation_message(total_rows: int, measures_in_actuals: list[str], metadata_measures: set[str]) -> Tuple[str, list[str]]:
    missing_indices = [idx for idx, value in enumerate(measures_in_actuals) if value and value not in metadata_measures]
    if not missing_indices:
        msg = (
            "Data import completed. Number of rows imported: {0}. Number of rows skipped: none. "
            "No mismatch found with existing metadata"
        ).format(total_rows)
        return msg, []

    skipped = len(missing_indices)
    imported = max(total_rows - skipped, 0)
    missing_value = measures_in_actuals[missing_indices[0]]
    msg = (
        "Data import completed with errors. Number of rows imported: {0}. Number of rows skipped: {1}. "
        "{2} member found in the load file is not present in the existing metadata in the system."
    ).format(imported, skipped, missing_value)
    return msg, [measures_in_actuals[idx] for idx in missing_indices]


def _generate_docx(cfg: Dict[str, str], message: str) -> Path:
    doc = Document()

    header = doc.add_paragraph()
    run = header.add_run(f"{cfg['admin_name']}, your validation details are ready.")
    run.font.bold = True
    run.font.size = Pt(10)

    doc.add_paragraph()
    details = doc.add_paragraph()
    run = details.add_run("Details")
    run.font.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0, 0, 128)

    doc.add_paragraph()
    table = doc.add_table(rows=0, cols=2)
    table.autofit = True
    body_lines = [
        ("Job Name:", "Load Data"),
        ("Job Type:", "Import Data"),
        ("Username:", cfg["username"]),
        ("Status Details:", message),
    ]
    for label, value in body_lines:
        row_cells = table.add_row().cells
        label_run = row_cells[0].paragraphs[0].add_run(label)
        label_run.font.bold = True
        label_run.font.size = Pt(8.5)
        label_run.font.color.rgb = RGBColor(128, 128, 128)
        value_run = row_cells[1].paragraphs[0].add_run(value)
        value_run.font.size = Pt(8.5)
        value_run.font.color.rgb = RGBColor(0, 0, 0)

    temp_dir = Path(tempfile.mkdtemp(prefix="validation_docx_"))
    docx_path = temp_dir / f"Validation_Report_{int(time.time())}.docx"
    doc.save(str(docx_path))
    return docx_path


def _upload_attachment(cfg: Dict[str, str], file_path: Path) -> str:
    remote_name = file_path.name
    encoded_name = quote(remote_name, safe="")
    upload_url = (
        f"{cfg['base_url']}/interop/rest/{cfg['upload_version']}"
        f"/applicationsnapshots/{encoded_name}/contents"
    )

    with open(file_path, "rb") as file_body:
        raw_bytes = file_body.read()

    auth_header = "Basic " + base64.b64encode(f"{cfg['username']}:{cfg['password']}".encode()).decode()
    headers = {
        "Authorization": auth_header,
        "Content-Type": "application/octet-stream",
        "Content-Length": str(len(raw_bytes)),
    }

    resp = requests.post(
        upload_url,
        headers=headers,
        params={"isForce": "true"},
        data=raw_bytes,
        timeout=120,
    )
    resp.raise_for_status()
    return remote_name


def _send_notification(cfg: Dict[str, str], message: str) -> None:
    attachment_path = _generate_docx(cfg, message)
    remote_name = _upload_attachment(cfg, attachment_path)
    payload = {
        "subject": f"{cfg['admin_name']}, your data validation is ready.",
        "body": f"{cfg['admin_name']}, your data validation is ready. Please find the validation output attached.",
        "to": cfg["admin_email"],
        "parameters": {
            "attachments": remote_name,
        },
    }
    email_url = f"{cfg['base_url']}/interop/rest/v2/mails/send"
    resp = requests.post(
        email_url,
        auth=(cfg["username"], cfg["password"]),
        headers={"Content-Type": "application/json"},
        data=json.dumps(payload),
        timeout=60,
    )
    if not resp.ok:
        raise RuntimeError(f"Send mail failed: {resp.status_code} {resp.text}")
    print("Send mail response:", resp.text)


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate measures and send result email")
    parser.add_argument("--entity", default="LE_ANZ", help="Entity identifier (default: LE_ANZ)")
    parser.add_argument(
        "--dataset",
        choices=["Actuals", "Forecast"],
        default="Actuals",
        help="Dataset suffix (default: Actuals)",
    )
    parser.add_argument(
        "--send-email",
        dest="send_email",
        action="store_true",
        help="Send validation result email (enabled by default; requires ADMIN_EMAIL_ID, ADMIN_NAME, EPM_UPLOAD_VERSION).",
    )
    parser.add_argument(
        "--no-send-email",
        dest="send_email",
        action="store_false",
        help="Disable validation result email.",
    )
    parser.set_defaults(send_email=True)
    args = parser.parse_args()

    cfg = _load_config()
    local_folder = Path(cfg["local_folder"])

    try:
        data_csv = local_folder / f"{args.entity} {args.dataset}.csv"
        if not data_csv.exists():
            raise FileNotFoundError(f"Data file not found: {data_csv}")

        try:
            export_zip = _find_latest_export_zip(local_folder)
            measures_csv = _extract_measures_csv(export_zip)
        except (FileNotFoundError, zipfile.BadZipFile):
            # Fallback to already exported local measures CSV.
            measures_csv = _find_local_measures_csv(local_folder)

        total_rows, actual_measures = _load_actual_measures(data_csv)
        metadata_measures = _load_measures(measures_csv)

        validation_details, _missing = _build_validation_message(total_rows, actual_measures, metadata_measures)
        print(validation_details)

        if args.send_email:
            missing_email_cfg = [
                key
                for key in ("admin_email", "admin_name", "upload_version")
                if not cfg.get(key)
            ]
            if missing_email_cfg:
                print(
                    "Email notification skipped due to missing environment variables: "
                    + ", ".join(missing_email_cfg)
                )
            else:
                _send_notification(cfg, validation_details)
                print("Notification sent successfully.")
        else:
            print("Email notification skipped (use --send-email to enable).")
        return 0
    finally:
        # Cleanup temp dirs created during extraction
        for temp_dir in Path(tempfile.gettempdir()).glob("validation_zip_*"):
            shutil.rmtree(temp_dir, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())