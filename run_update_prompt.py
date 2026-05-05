#!/usr/bin/env python
"""Prompt-to-metadata wrapper.

Parses natural-language member requests, writes rows into:
- AccountsToBeAdded.xlsx (Account/Measures)
- LOBsToBeAdded.xlsx (LOB)

Then invokes update_metadata.py to run export/import metadata jobs.
"""

import argparse
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

PROJECT_DIR = Path(__file__).resolve().parent
ACCOUNT_XLSX = PROJECT_DIR / "AccountsToBeAdded.xlsx"
LOB_XLSX = PROJECT_DIR / "LOBsToBeAdded.xlsx"


def _ensure_workbook(path: Path) -> None:
    if path.exists():
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("Sheet1")
    ws.append(["Name", "Parent", "Description"])
    wb.save(path)


def _upsert_row(path: Path, member: str, parent: str, desc: str) -> str:
    _ensure_workbook(path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("Sheet1")

    # find exact existing row
    for row in ws.iter_rows(min_row=2, max_col=3):
        name = str(row[0].value or "").strip()
        par = str(row[1].value or "").strip()
        dsc = str(row[2].value or "").strip()
        if name.lower() == member.lower() and par.lower() == parent.lower() and dsc.lower() == desc.lower():
            wb.close()
            return "exists"

    ws.append([member, parent, desc])
    wb.save(path)
    wb.close()
    return "added"


def _parse_prompt(prompt: str) -> list[tuple[str, str, str, str]]:
    parts = [p.strip() for p in re.split(r";|\n", prompt) if p.strip()]
    out: list[tuple[str, str, str, str]] = []
    pat = re.compile(
        r"add\s+(?:member\s+)?(?P<member>[^,]+),\s*with\s+(?P<parent>[^,]+?)\s+parent,\s*with\s+(?P<desc>[^,]+?)\s+description\s+in\s+the\s+(?P<dim>lob|account)\s+dimension",
        re.IGNORECASE,
    )
    for part in parts:
        m = pat.search(part)
        if not m:
            continue
        out.append(
            (
                re.sub(r"\s+member$", "", m.group("member").strip(), flags=re.IGNORECASE),
                m.group("parent").strip(),
                m.group("desc").strip(),
                m.group("dim").strip().lower(),
            )
        )
    return out


def _invoke_update(prompt: str | None = None, prompt_file: Path | None = None) -> int:
    text = ""
    if prompt:
        text = prompt
    elif prompt_file:
        text = prompt_file.read_text(encoding="utf-8", errors="ignore")
    else:
        return 0

    requests = _parse_prompt(text)
    if not requests:
        print("No valid member requests found in prompt. Nothing to update.")
        return 0

    added_any = False
    for member, parent, desc, dim in requests:
        target = LOB_XLSX if dim == "lob" else ACCOUNT_XLSX
        status = _upsert_row(target, member, parent, desc)
        if status == "exists":
            print(f"Member already exists in workbook: {member} ({dim.upper()})")
        else:
            added_any = True
            print(f"Added member to workbook: {member} ({dim.upper()})")

    if not added_any:
        print("Member already exists. Stopping metadata update.")
        return 0

    cmd = [sys.executable, str(PROJECT_DIR / "update_metadata.py")]
    result = subprocess.run(cmd, check=False, cwd=str(PROJECT_DIR))
    return result.returncode


def main() -> int:
    parser = argparse.ArgumentParser(description="Wrapper to trigger metadata update via prompt.")
    parser.add_argument("--prompt", help="Natural‑language prompt string.")
    parser.add_argument(
        "--prompt-file",
        type=Path,
        help="Path to a file containing one or more prompts (one per line).",
    )
    args = parser.parse_args()

    # If both are supplied, prefer the explicit ``--prompt`` value.
    if args.prompt:
        return _invoke_update(prompt=args.prompt)
    if args.prompt_file:
        return _invoke_update(prompt_file=args.prompt_file)
    # No prompt supplied – per requirement, ignore the prompt section.
    print("No prompt provided; exiting without running update.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
